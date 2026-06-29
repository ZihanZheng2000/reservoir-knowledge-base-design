from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.util import Inches, Pt


OUT = Path(
    r"D:\Auto Research\case study\reservoir-knowledge-base-agentic-workflow"
    r"\runs\lake_powell_20260626\synthesis_stage3_rerun_20260626_two_step"
    r"\lake_powell_method_meeting_slides.pptx"
)

NOTES_OUT = OUT.with_suffix(".talking_points.md")


COLORS = {
    "ink": RGBColor(34, 42, 53),
    "muted": RGBColor(92, 102, 117),
    "blue": RGBColor(42, 95, 170),
    "teal": RGBColor(26, 137, 125),
    "amber": RGBColor(197, 122, 36),
    "red": RGBColor(176, 70, 70),
    "line": RGBColor(214, 220, 229),
    "soft": RGBColor(244, 247, 250),
    "white": RGBColor(255, 255, 255),
}


def add_textbox(slide, x, y, w, h, text, size=18, bold=False, color="ink", align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    p = tf.paragraphs[0]
    if align:
        p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = COLORS[color]
    run.font.name = "Aptos"
    return box


def add_title(slide, title, subtitle=None):
    add_textbox(slide, 0.55, 0.32, 12.2, 0.55, title, size=25, bold=True, color="ink")
    if subtitle:
        add_textbox(slide, 0.58, 0.86, 11.8, 0.36, subtitle, size=11, color="muted")
    line = slide.shapes.add_shape(1, Inches(0.55), Inches(1.18), Inches(12.25), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = COLORS["line"]
    line.line.color.rgb = COLORS["line"]


def add_bullets(slide, x, y, w, h, bullets, size=15, color="ink", leading=1.08):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    for idx, item in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(size)
        p.font.name = "Aptos"
        p.font.color.rgb = COLORS[color]
        p.space_after = Pt(5 * leading)
    return box


def add_card(slide, x, y, w, h, title, body, accent="blue", title_size=14, body_size=12):
    shape = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLORS["white"]
    shape.line.color.rgb = COLORS["line"]
    stripe = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(0.08), Inches(h))
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = COLORS[accent]
    stripe.line.color.rgb = COLORS[accent]
    add_textbox(slide, x + 0.22, y + 0.14, w - 0.34, 0.34, title, size=title_size, bold=True, color=accent)
    add_textbox(slide, x + 0.22, y + 0.55, w - 0.34, h - 0.68, body, size=body_size, color="ink")


def add_footer(slide, n):
    add_textbox(slide, 11.9, 7.05, 0.75, 0.22, str(n), size=9, color="muted", align=PP_ALIGN.RIGHT)


def add_table(slide, x, y, w, h, rows, cols, data, col_widths=None, font_size=10):
    table_shape = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(h))
    table = table_shape.table
    if col_widths:
        for i, cw in enumerate(col_widths):
            table.columns[i].width = Inches(cw)
    for r in range(rows):
        for c in range(cols):
            cell = table.cell(r, c)
            cell.text = data[r][c]
            cell.margin_left = Inches(0.06)
            cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.04)
            cell.margin_bottom = Inches(0.04)
            fill = COLORS["soft"] if r == 0 else COLORS["white"]
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill
            cell.text_frame.word_wrap = True
            for p in cell.text_frame.paragraphs:
                p.font.name = "Aptos"
                p.font.size = Pt(font_size if r else font_size + 0.5)
                p.font.bold = r == 0
                p.font.color.rgb = COLORS["ink"]
    return table_shape


def add_flow(slide, y, labels, colors=None):
    colors = colors or ["blue"] * len(labels)
    x = 0.85
    box_w = 2.12
    gap = 0.35
    for i, label in enumerate(labels):
        shape = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(box_w), Inches(0.7))
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLORS[colors[i]]
        shape.line.color.rgb = COLORS[colors[i]]
        add_textbox(slide, x + 0.06, y + 0.18, box_w - 0.12, 0.25, label, size=12, bold=True, color="white", align=PP_ALIGN.CENTER)
        if i < len(labels) - 1:
            add_textbox(slide, x + box_w + 0.07, y + 0.22, 0.22, 0.2, "→", size=20, bold=True, color="muted", align=PP_ALIGN.CENTER)
        x += box_w + gap


def build_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    slides_notes: list[tuple[str, list[str]]] = []

    def slide(title, subtitle=None):
        s = prs.slides.add_slide(blank)
        add_title(s, title, subtitle)
        add_footer(s, len(prs.slides))
        return s

    # 1
    s = prs.slides.add_slide(blank)
    add_textbox(s, 0.75, 0.68, 11.7, 0.65, "Explaining the Revised Lake Powell Synthesis Method", 28, True, "ink")
    add_textbox(s, 0.78, 1.42, 10.8, 0.42, "Meeting focus: why CKP matters, how false discrepancies are reduced, and how the agent workflow is verified and improved", 15, False, "muted")
    add_flow(s, 2.45, ["Source docs", "KU", "CKP", "Synthesis"], ["blue", "teal", "amber", "blue"])
    add_card(s, 0.9, 4.0, 3.7, 1.5, "Core claim", "The revised workflow separates evidence preservation, knowledge organization, and cross-document reasoning.", "blue")
    add_card(s, 4.85, 4.0, 3.7, 1.5, "Key change", "CKP is added as the intermediate operational-context layer before synthesis.", "amber")
    add_card(s, 8.8, 4.0, 3.7, 1.5, "Main result", "Fewer false conflicts, clearer synthesis types, and stronger evidence traceability.", "teal")
    add_footer(s, 1)
    slides_notes.append(("Slide 1", ["Open by saying this is a method explanation, not just a result summary.", "The deck answers Prof. Cai's questions directly."]))

    # 2
    s = slide("Three-Layer Method Framework")
    add_card(s, 0.75, 1.55, 3.8, 4.7, "Layer 1: Workflow", "Five-stage process that turns reservoir documents into source-grounded KUs, CKPs, synthesis cards, and report/retrieval outputs.", "blue", 15, 13)
    add_card(s, 4.78, 1.55, 3.8, 4.7, "Layer 2: Agent Architecture", "Project docs, skills, templates, scripts, run folders, and human gates constrain the AI agent so the workflow is repeatable.", "teal", 15, 13)
    add_card(s, 8.82, 1.55, 3.8, 4.7, "Layer 3: Debugging", "Review outputs, identify failure modes, revise prompts/schemas/workflow rules, rerun affected stages, and compare quality.", "amber", 15, 13)
    slides_notes.append(("Slide 2", ["Use this as the roadmap for the meeting.", "Emphasize that the agent is embedded in a method, not free-form summarization."]))

    # 3
    s = slide("Layer 1: Five-Step Knowledge-Base Workflow")
    data = [
        ["Step", "Purpose", "Input", "Output"],
        ["1. Source acquisition", "Preserve and screen operation-relevant sources", "URLs, PDFs, web pages", "Inventory, raw files, extracted text"],
        ["2. KU extraction", "Extract source-grounded operational facts", "Source chunks/text", "Operational KUs with locators"],
        ["3. CKP consolidation", "Organize KUs into operational topics", "Validated KUs", "Consolidated Knowledge Points"],
        ["4. Synthesis", "Identify cross-KU relationships", "CKPs, KUs, locators", "Synthesis cards"],
        ["5. Report / retrieval", "Make outputs usable and reviewable", "KUs + synthesis", "Reports, retrieval records, claim maps"],
    ]
    add_table(s, 0.55, 1.45, 12.25, 4.95, 6, 4, data, [2.0, 4.0, 2.85, 3.4], 9.2)
    add_textbox(s, 0.8, 6.58, 11.8, 0.35, "Key point: CKP is the new intermediate layer that creates operational context before synthesis.", 13, True, "amber")
    slides_notes.append(("Slide 3", ["Walk through the five steps briefly.", "Say the rest of the deck focuses on why step 3 changes the quality of step 4."]))

    # 4
    s = slide("Step 1: Source Acquisition")
    add_card(s, 0.75, 1.45, 3.65, 4.55, "Purpose", "Build a preserved and reviewable source corpus before any knowledge extraction.", "blue")
    add_card(s, 4.65, 1.45, 3.65, 4.55, "How", "Discover candidate sources, screen by operation relevance and source tier, preserve raw files, extract text, and validate inventory.", "teal")
    add_card(s, 8.55, 1.45, 3.65, 4.55, "Input / Output", "Input: candidate URLs, PDFs, reports, web pages.\n\nOutput: source_inventory.jsonl, sources/raw/, sources/text/, metadata.", "amber")
    slides_notes.append(("Slide 4", ["Purpose is evidence preservation before interpretation.", "This is important because later claims can only be as good as the source corpus."]))

    # 5
    s = slide("Step 2: KU Extraction")
    add_card(s, 0.7, 1.4, 3.8, 4.9, "Purpose", "Convert documents into small, source-grounded operational facts.", "blue")
    add_card(s, 4.75, 1.4, 3.8, 4.9, "How", "Parse source text into chunks/sections, extract document-level KUs, and keep source ID, locator, evidence quote, engineering dimension, and relevance.", "teal")
    add_card(s, 8.8, 1.4, 3.8, 4.9, "Example", "A KU records that WY2026 release was projected as 7.48 maf under the Mid-Elevation Release Tier, with source URL and chunk locator.", "amber")
    add_textbox(s, 0.85, 6.55, 11.6, 0.34, "KU preserves evidence before interpretation.", 13, True, "blue", PP_ALIGN.CENTER)
    slides_notes.append(("Slide 5", ["Explain KU as the evidence-preserving unit.", "Each KU should be traceable back to source evidence."]))

    # 6
    s = slide("Step 3: CKP Consolidation: The Key Intermediate Layer")
    add_textbox(s, 0.75, 1.45, 12.0, 0.45, "CKP is not just another summary. It is a semantic organization layer before synthesis.", 16, True, "amber")
    add_bullets(s, 0.9, 2.15, 5.7, 3.65, [
        "Groups fragmented KUs into stable reservoir-operation topics.",
        "Preserves KU IDs and source coverage.",
        "Keeps values, rules, thresholds, scenarios, dates, and decision stages together.",
        "Gives synthesis the context needed to reason correctly.",
    ], 15)
    add_card(s, 7.0, 2.0, 5.2, 3.6, "Approved CKP dimensions", "Operation Rules\nInfrastructure\nReal-Time / Emergency Operations\nCoordination / Governance\nOperational Data\nResearch / Modeling / Analysis\nEvidence Gap / Uncertainty", "teal", 14, 12)
    add_textbox(s, 0.9, 6.35, 11.4, 0.42, "Core meaning: CKP transforms scattered source-level facts into organized operational context.", 14, True, "amber", PP_ALIGN.CENTER)
    slides_notes.append(("Slide 6", ["This is the slide to slow down on.", "Say CKP creates the missing context between KU and synthesis."]))

    # 7
    s = slide("Example: CKP Reduces False Source Discrepancies")
    add_card(s, 0.65, 1.45, 3.75, 4.85, "Before CKP", "Separate KUs mention 8.23 maf, 7.48 maf, 6.00 maf, CRSS logic, and 24-Month Study scenarios.\n\nA naive synthesis may label them as conflicting release values.", "red", 14, 12)
    add_card(s, 4.8, 1.45, 3.75, 4.85, "CKP grouping", "These KUs are grouped under Operation Rules:\n\nrelease tiers, objective releases, drought-response reductions, and model-based operating decisions.", "amber", 14, 12)
    add_card(s, 8.95, 1.45, 3.75, 4.85, "After CKP", "The synthesis can interpret them as different rule layers, scenarios, and decision stages rather than a proven source conflict.", "teal", 14, 12)
    add_textbox(s, 1.0, 6.55, 11.2, 0.34, "Specific result: reduced false source discrepancies.", 13, True, "teal", PP_ALIGN.CENTER)
    slides_notes.append(("Slide 7", ["Use this as the main concrete example.", "Say 7.48 and 6.00 maf are better understood as status/timing differences."]))

    # 8
    s = slide("Step 4: Cross-KU Synthesis")
    add_card(s, 0.75, 1.45, 3.7, 4.65, "Purpose", "Identify relationships across organized CKPs and supporting KUs.", "blue")
    add_card(s, 4.85, 1.45, 3.7, 4.65, "How", "Use a prompt-defined synthesis taxonomy and cite supporting KU IDs for each card.", "teal")
    add_card(s, 8.95, 1.45, 3.7, 4.65, "Output", "Synthesis cards with analysis_type, summary, interpretation, evidence depth, KU IDs, locator summary, and next step.", "amber")
    add_flow(s, 6.32, ["CKPs", "Candidate relationship", "Source check", "Synthesis card"], ["amber", "blue", "teal", "blue"])
    slides_notes.append(("Slide 8", ["Synthesis is not generic summary.", "It classifies evidence relationships."]))

    # 9
    s = slide("Improved Classification Comes from Prompt-Level Topic Definitions")
    add_card(s, 0.65, 1.45, 3.75, 4.85, "source_discrepancy", "Only when sources disagree about the same object under the same status, date, scenario, and decision context.", "red", 13, 11.5)
    add_card(s, 4.8, 1.45, 3.75, 4.85, "complementary_evidence", "When different sources provide different parts of the same decision chain, such as forecast, model, rule, and action.", "blue", 13, 11.5)
    add_card(s, 8.95, 1.45, 3.75, 4.85, "operational_tradeoff", "When reservoir objectives pull decisions in different directions, such as storage protection versus downstream release or hydropower.", "teal", 13, 11.5)
    add_textbox(s, 0.85, 6.52, 11.6, 0.38, "The classification improved because the synthesis prompt/protocol redefined each topic boundary.", 13, True, "ink", PP_ALIGN.CENTER)
    slides_notes.append(("Slide 9", ["Say this improvement is mainly prompt/protocol design.", "It teaches the model what each synthesis type means."]))

    # 10
    s = slide("Step 5: Report and Retrieval Outputs")
    add_bullets(s, 0.9, 1.55, 5.9, 4.7, [
        "Generate reports only from validated KUs and synthesis cards.",
        "Preserve claim-evidence maps for review.",
        "Prepare retrieval records and benchmark questions.",
        "Keep final narrative connected to the knowledge base.",
    ], 16)
    add_card(s, 7.1, 1.7, 5.15, 3.9, "Why this matters", "The final output is not just a readable report. It remains connected to source IDs, KU IDs, synthesis IDs, and evidence locators.", "blue", 15, 13)
    slides_notes.append(("Slide 10", ["This closes Layer 1.", "Emphasize that final outputs remain auditable."]))

    # 11
    s = slide("Layer 2: AI Agent Architecture")
    data = [
        ["Component", "Role"],
        ["Project docs", "Define principles, stage boundaries, required artifacts, and human gates."],
        ["Skills", "Stage-specific instructions for source acquisition, KU extraction, synthesis, retrieval, and reporting."],
        ["Templates", "Standardize manifests, validation summaries, and report structures."],
        ["Scripts", "Run deterministic checks for schemas, references, source inventory, and synthesis structure."],
        ["Run folders", "Version outputs and prevent overwriting previous runs."],
        ["Human gates", "Require approval for thresholds, schemas, report-ready claims, and dataset release."],
    ]
    add_table(s, 0.75, 1.35, 11.85, 5.65, 7, 2, data, [2.65, 9.2], 10.4)
    slides_notes.append(("Slide 11", ["Say the agent works inside a controlled architecture.", "This makes the workflow repeatable rather than ad hoc."]))

    # 12
    s = slide("How the Agent Keeps the System Working")
    add_card(s, 0.7, 1.4, 3.75, 4.85, "Agent coordination", "The agent selects the relevant skill, follows project rules, runs scripts, reads outputs, and decides whether a stage should be rerun.", "blue", 14, 12)
    add_card(s, 4.8, 1.4, 3.75, 4.85, "Deterministic checks", "Python scripts validate required fields, allowed analysis types, KU references, locator metadata, and output consistency.", "teal", 14, 12)
    add_card(s, 8.9, 1.4, 3.75, 4.85, "Human control", "The workflow keeps gates for source thresholds, schema changes, report-ready synthesis claims, and release decisions.", "amber", 14, 12)
    slides_notes.append(("Slide 12", ["This slide answers why the platform can execute the method.", "The agent is a coordinator plus checker, not just a writer."]))

    # 13
    s = slide("Source Verification and Traceability Mechanism")
    add_flow(s, 1.55, ["Synthesis card", "KU IDs", "Source locator", "Evidence text"], ["blue", "teal", "amber", "blue"])
    add_card(s, 0.75, 2.95, 3.7, 3.25, "Recorded fields", "based_on_ku_ids\nevidence_depth\nsource_locator_summary\nsource_verification_note", "blue", 14, 12)
    add_card(s, 4.85, 2.95, 3.7, 3.25, "What code checks", "Required fields, valid analysis type, cited KU IDs exist, source-checked cards include locator summary and verification note.", "teal", 14, 12)
    add_card(s, 8.95, 2.95, 3.7, 3.25, "Important nuance", "The script validates traceability structure. Source rereading is enforced by the workflow/prompt and recorded in verification metadata.", "amber", 14, 12)
    add_textbox(s, 0.95, 6.55, 11.3, 0.35, "This is a semi-automated verification gate: structured validation plus source rereading.", 13, True, "teal", PP_ALIGN.CENTER)
    slides_notes.append(("Slide 13", ["Be careful not to overclaim full automatic fact checking.", "Say the code verifies the traceability structure."]))

    # 14
    s = slide("Layer 3: Debugging and Refinement")
    add_flow(s, 1.6, ["Run", "Inspect", "Diagnose", "Revise", "Rerun"], ["blue", "teal", "amber", "blue", "teal"])
    add_bullets(s, 1.0, 3.0, 5.5, 3.1, [
        "Review generated KUs, CKPs, and synthesis cards.",
        "Identify failure modes from actual outputs.",
        "Revise prompt taxonomy, schemas, or workflow rules.",
        "Rerun affected stages first, then full workflow if needed.",
    ], 15)
    add_card(s, 7.0, 3.0, 5.15, 2.6, "Goal", "Make the system improve through output diagnosis and rerun, instead of relying on a single prompt attempt.", "amber", 15, 13)
    slides_notes.append(("Slide 14", ["This is the third layer: how the method improves.", "It is an iterative workflow, not one-shot generation."]))

    # 15
    s = slide("Debugging Example: False Discrepancy")
    add_card(s, 0.65, 1.4, 3.75, 4.9, "Observed issue", "Different release values could be treated as source conflicts.", "red", 14, 12)
    add_card(s, 4.8, 1.4, 3.75, 4.9, "Diagnosis", "The system needed a layer to distinguish objective release, projected release, adjusted release, scenario value, and model logic.", "amber", 14, 12)
    add_card(s, 8.95, 1.4, 3.75, 4.9, "Revision and result", "Add CKP consolidation before synthesis. Release values are grouped under Operation Rules, so the final card treats them as status/context differences.", "teal", 14, 12)
    slides_notes.append(("Slide 15", ["Tie this back to Prof. Cai's question.", "This is the concrete effect of CKP."]))

    # 16
    s = slide("Debugging Example: Synthesis Classification")
    add_card(s, 0.7, 1.4, 3.8, 4.9, "Observed issue", "Synthesis categories could be too broad or mixed.", "red", 14, 12)
    add_card(s, 4.75, 1.4, 3.8, 4.9, "Revision", "Redefine prompt-level taxonomy and add decision rules for each analysis type.", "blue", 14, 12)
    add_card(s, 8.8, 1.4, 3.8, 4.9, "Improved outputs", "Release values: complementary/status metadata.\nStorage vs delivery: operational tradeoff.\nUncertainty-to-trigger logic: evidence gap.\nPost-2026 coordination: outstanding issue.", "teal", 14, 11.5)
    slides_notes.append(("Slide 16", ["Say the classification improvement was a prompt/protocol improvement.", "Examples show better alignment with reservoir operation reasoning."]))

    # 17
    s = slide("How We Know It Improved")
    add_bullets(s, 0.85, 1.45, 5.75, 4.85, [
        "CKPs make reusable operational topics visible.",
        "Synthesis cards cite multiple KU IDs.",
        "Source-checked cards include locator summaries and verification notes.",
        "Categories better match reservoir-operation logic.",
        "No forced source discrepancy when no true same-object/same-status conflict is established.",
    ], 15)
    add_card(s, 7.0, 1.7, 5.15, 3.75, "Lake Powell revised run", "8 synthesis cards:\n2 recurring findings\n2 complementary evidence\n2 operational tradeoffs\n1 evidence gap\n1 outstanding issue\n\nValidation script: zero structural errors.", "teal", 15, 12.5)
    slides_notes.append(("Slide 17", ["Mention this is output-level evidence, not final domain validation.", "It shows the workflow behaves better structurally and semantically."]))

    # 18
    s = slide("Takeaway")
    add_textbox(s, 0.85, 1.65, 11.6, 0.55, "The revised workflow works because it separates evidence preservation, knowledge organization, and cross-document reasoning.", 20, True, "ink", PP_ALIGN.CENTER)
    add_card(s, 1.0, 2.65, 3.45, 3.1, "1. CKP", "Provides operational context before synthesis.", "amber", 16, 14)
    add_card(s, 4.95, 2.65, 3.45, 3.1, "2. Taxonomy", "Improves synthesis classification through clearer prompt-level topic definitions.", "blue", 16, 14)
    add_card(s, 8.9, 2.65, 3.45, 3.1, "3. Traceability", "Keeps high-level claims connected to KU IDs, source locators, and verification notes.", "teal", 16, 14)
    add_textbox(s, 0.95, 6.4, 11.4, 0.38, "Overall: heterogeneous reservoir documents become a structured, auditable, and improvable operational knowledge base.", 14, True, "blue", PP_ALIGN.CENTER)
    slides_notes.append(("Slide 18", ["Close by returning to Prof. Cai's questions.", "CKP is the main methodological answer."]))

    prs.save(OUT)
    notes = ["# Lake Powell Method Meeting Slides - Talking Points", ""]
    for title, bullets in slides_notes:
        notes.append(f"## {title}")
        for b in bullets:
            notes.append(f"- {b}")
        notes.append("")
    NOTES_OUT.write_text("\n".join(notes), encoding="utf-8")


if __name__ == "__main__":
    build_deck()
